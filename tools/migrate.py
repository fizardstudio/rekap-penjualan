"""Migrasi Excel lama -> seed JSON untuk aplikasi Rekap Penjualan.
Pemakaian: python3 migrate.py <folder Bahan> <output.json>
"""
import openpyxl, re, json, sys, datetime, collections, os

BAHAN = sys.argv[1]
OUT = sys.argv[2]

def P(name):
    return os.path.join(BAHAN, name)

def s(v):
    return v.strip() if isinstance(v, str) else v

def num(v):
    if v is None or v == "":
        return 0
    if isinstance(v, (int, float)):
        return v
    t = str(v).strip().replace("Rp", "").replace(" ", "")
    if re.fullmatch(r"[\d\.]+,\d+", t):
        t = t.replace(".", "").replace(",", ".")
    else:
        t = t.replace(",", "")
    try:
        return float(t)
    except ValueError:
        return 0

def clean_name(v):
    return re.sub(r"\s+", " ", str(v or "").replace("\t", " ")).strip()

# ---------------------------------------------------------------- master barang
def load_barang():
    ws = openpyxl.load_workbook(P("HARGA ATK INSTANSI Mei 2026-1.xlsx")).active
    out = []
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
        if not r[2]:
            continue
        out.append({
            "id": i,
            "kode": str(r[1]).strip() if r[1] is not None else "",
            "nama": clean_name(r[2]),
            "merk": clean_name(r[3]),
            "satuan": clean_name(r[4]).upper(),
            "harga": round(num(r[5])),
            "jenis": "B",
        })
    return out

# ---------------------------------------------------------------- master satker
KANTAH = [  # (sheet BELANJA, kode kantah, nama)
    ("KANWIL JABAR", "32", "KANTOR WILAYAH BPN PROV. JAWA BARAT"),
    ("KANWIL JATENG", "33", "KANTOR WILAYAH BPN PROV. JAWA TENGAH"),
    ("SURAKARTA", "33.72", "KANTOR PERTANAHAN KOTA SURAKARTA"),
    ("KUANSING", "14.09", "KANTOR PERTANAHAN KAB. KUANTAN SINGINGI"),
    ("KOTA BANDUNG", "32.73", "KANTOR PERTANAHAN KOTA BANDUNG"),
    ("KAB. SUMEDANG", "32.11", "KANTOR PERTANAHAN KAB. SUMEDANG"),
    ("KOTA CIMAHI", "32.77", "KANTOR PERTANAHAN KOTA CIMAHI"),
    ("KOTA BANJAR", "32.79", "KANTOR PERTANAHAN KOTA BANJAR"),
    ("KOTA SUKABUMI", "32.72", "KANTOR PERTANAHAN KOTA SUKABUMI"),
    ("KOTA DEPOK", "32.76", "KANTOR PERTANAHAN KOTA DEPOK"),
    ("KOTA BEKASI", "32.75", "KANTOR PERTANAHAN KOTA BEKASI"),
    ("KOTA BOGOR", "32.71", "KANTOR PERTANAHAN KOTA BOGOR"),
    ("KOTA TASIK", "32.78", "KANTOR PERTANAHAN KOTA TASIKMALAYA"),
    ("KOTA CIREBON", "32.74", "KANTOR PERTANAHAN KOTA CIREBON"),
    ("KAB. GARUT", "32.05", "KANTOR PERTANAHAN KAB. GARUT"),
    ("KAB. TASIKMALAYA", "32.06", "KANTOR PERTANAHAN KAB. TASIKMALAYA"),
    ("KAB. CIAMIS", "32.07", "KANTOR PERTANAHAN KAB. CIAMIS"),
    ("KAB. PANGANDARAN", "32.18", "KANTOR PERTANAHAN KAB. PANGANDARAN"),
    ("KAB. KUNINGAN", "32.08", "KANTOR PERTANAHAN KAB. KUNINGAN"),
    ("KAB. INDRAMAYU", "32.12", "KANTOR PERTANAHAN KAB. INDRAMAYU"),
    ("KAB. SUBANG", "32.13", "KANTOR PERTANAHAN KAB. SUBANG"),
    ("KAB. KARAWANG", "32.15", "KANTOR PERTANAHAN KAB. KARAWANG"),
    ("KAB. PURWAKARTA", "32.14", "KANTOR PERTANAHAN KAB. PURWAKARTA"),
    ("KAB. BANDUNG", "32.04", "KANTOR PERTANAHAN KAB. BANDUNG"),
    ("KBB", "32.17", "KANTOR PERTANAHAN KAB. BANDUNG BARAT"),
    ("KAB. BEKASI", "32.16", "KANTOR PERTANAHAN KAB. BEKASI"),
    ("KAB. BOGOR I", "32.01A", "KANTOR PERTANAHAN KAB. BOGOR I"),
    ("KAB. BOGOR II", "32.01B", "KANTOR PERTANAHAN KAB. BOGOR II"),
    ("KAB. MAJALENGKA", "32.10", "KANTOR PERTANAHAN KAB. MAJALENGKA"),
    ("KAB. SUKABUMI", "32.02", "KANTOR PERTANAHAN KAB. SUKABUMI"),
    ("KAB. CIREBON", "32.09", "KANTOR PERTANAHAN KAB. CIREBON"),
    ("CIANJUR", "32.03", "KANTOR PERTANAHAN KAB. CIANJUR"),
]

# sheet rekap Digipay -> (nama satker, sheet BELANJA bila BPN)
DIGI_SHEET = {
    "BPN KAB. BEKASI": "KAB. BEKASI", "BPN MAJALENGKA": "KAB. MAJALENGKA",
    "BPN KAB. BOGOR": "KAB. BOGOR I", "BPN SUBANG": "KAB. SUBANG", "BPN GARUT": "KAB. GARUT",
    "BPN KARAWANG": "KAB. KARAWANG", "BPN KBB": "KBB", "BPN KOTA TASIKMALAYA": "KOTA TASIK",
    "BPN KOTA CIMAHI": "KOTA CIMAHI", "BPN BANDUNG": "KOTA BANDUNG", "BPN PURWAKARTA": "KAB. PURWAKARTA",
    "BPN KAB. TASIK": "KAB. TASIKMALAYA", "BPN CIANJUR": "CIANJUR", "BPN CIAMIS": "KAB. CIAMIS",
    "BPN BANJAR": "KOTA BANJAR", "BPN KUNINGAN": "KAB. KUNINGAN", "BPN KAB.SUKABUMI": "KAB. SUKABUMI",
    "BPN SUMEDANG": "KAB. SUMEDANG", "BPN KOTA SUKABUMI": "KOTA SUKABUMI",
}
NONBPN = {
    "BPS KOTA BOGOR": ("BADAN PUSAT STATISTIK KOTA BOGOR", "18740"),
    "BHLK": ("BALAI HIDROLIKA DAN LINGKUNGAN KEAIRAN", "420110"),
    "BHGK BANDUNG": ("BALAI HIDROLIKA DAN GEOTEKNIK KEAIRAN", "420108"),
    "BALAI AIR TANAH": ("BALAI AIR TANAH", "420109"),
    "ISBI BDG": ("INSTITUT SENI BUDAYA INDONESIA BANDUNG", ""),
    "BPS CIMAHI": ("BADAN PUSAT STATISTIK KOTA CIMAHI", "652229"),
    "DISDUK JABAR": ("DINAS PENDIDIKAN PROV. JAWA BARAT", ""),
    "B4T BDG": ("BALAI BESAR BAHAN DAN BARANG TEKNIK (B4T)", ""),
    "BALAI TEXTILE": ("BALAI BESAR STANDARDISASI DAN PELAYANAN JASA INDUSTRI TEKSTIL", ""),
}

def load_satker(wb_digi):
    # kode satker digipay dari HOME
    home = wb_digi["HOME"]
    kode_by_sheet = {}
    for r in home.iter_rows(min_row=6, max_row=40):
        kode, f = r[2].value, r[4].value
        if kode and isinstance(f, str):
            m = re.match(r"='?([^'!]+)'?!", f)
            if m:
                kode_by_sheet[m.group(1)] = str(kode)
    satker = []
    by_kantah_sheet, by_digi_sheet = {}, {}
    nid = 1
    for sheet, kk, nama in KANTAH:
        rec = {"id": nid, "nama": nama, "singkat": sheet, "grup": "BPN", "kodeKantah": kk,
               "kodeSatker": "", "wapu": True}
        satker.append(rec); by_kantah_sheet[sheet] = rec; nid += 1
    for ds, ks in DIGI_SHEET.items():
        rec = by_kantah_sheet[ks]
        if kode_by_sheet.get(ds):
            rec["kodeSatker"] = kode_by_sheet[ds]
        by_digi_sheet[ds] = rec
    for ds, (nama, kode) in NONBPN.items():
        rec = {"id": nid, "nama": nama, "singkat": ds, "grup": "Lainnya", "kodeKantah": "",
               "kodeSatker": kode_by_sheet.get(ds, kode), "wapu": True}
        satker.append(rec); by_digi_sheet[ds] = rec; nid += 1
    # Kode satker Kab. Tasikmalaya sudah ada dari HOME (429740). Garut dari video: 429734.
    return satker, by_kantah_sheet, by_digi_sheet

# ---------------------------------------------------------------- parser blok
def find_blocks(ws, vs=None):
    """Kembalikan daftar blok. ws = sheet rumus, vs = sheet nilai (data_only)."""
    vs = vs or ws
    blocks = []
    maxr = ws.max_row
    for row in ws.iter_rows():
        for c in row:
            if s(c.value) == "NO" and s(ws.cell(c.row, c.column + 1).value) in ("NAMA BARANG", "NAMA ITEM"):
                blocks.append({"hdr": c.row, "col": c.column})
                break
    for i, b in enumerate(blocks):
        end = blocks[i + 1]["hdr"] - 1 if i + 1 < len(blocks) else maxr
        b["end"] = end
        col = b["col"]
        items, r = [], b["hdr"] + 1
        total_rows = []
        while r <= end:
            nama = ws.cell(r, col + 1).value
            lab_f = s(ws.cell(r, col + 4).value)
            gf = ws.cell(r, col + 5).value
            gf = gf.upper().replace(" ", "") if isinstance(gf, str) else ""
            is_total = (isinstance(nama, str) and re.search(r"TOTAL|^DPP$|^PPN", nama.strip(), re.I)) or \
                       (isinstance(lab_f, str) and re.search(r"TOTAL|PPN|PPH|PEMBAYARAN|DPP", lab_f, re.I)) or \
                       gf.startswith("=SUM(") or "%" in gf
            if is_total:
                lab = (lab_f if isinstance(lab_f, str) else "") or (nama if isinstance(nama, str) else "")
                if not lab and "%" in gf:
                    lab = "PPN"
                total_rows.append((r, lab.upper(), num(vs.cell(r, col + 5).value)))
                r += 1
                continue
            if total_rows:
                break
            q = num(vs.cell(r, col + 2).value)
            h = num(vs.cell(r, col + 4).value)
            g = vs.cell(r, col + 5).value
            g = num(g) if isinstance(g, (int, float)) else q * h
            if q and abs(g - q * h) > 1:
                h = g / q
            if nama not in (None, "") or q:
                items.append({
                    "row": r, "nama": clean_name(nama), "qty": q,
                    "satuan": clean_name(ws.cell(r, col + 3).value).upper(),
                    "harga": h,
                })
            r += 1
        b["items"] = [it for it in items if it["nama"] and it["qty"]]
        b["totals"] = total_rows
        b["hasPPN"] = any("PPN" in t and v for _, t, v in total_rows)
        b["hasPPh"] = any("PPH" in t for _, t, _ in total_rows)
        b["exTotal"] = total_rows[-1][2] if total_rows else 0
        b["label"] = clean_name(ws.cell(b["hdr"] - 1, col + 1).value)
    return blocks

def date_from_text(t, year=None):
    if isinstance(t, datetime.datetime):
        return t.date()
    if not t:
        return None
    m = re.search(r"(\d{1,2})[-/\.](\d{1,2})[-/\.](\d{4})", str(t))
    if m:
        try:
            return datetime.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        except ValueError:
            return None
    return None

def mk_items(b, barang_idx):
    out = []
    for it in b["items"]:
        nm = it["nama"]
        base = re.sub(r"\s*-?\s*(OSI|KJS)\s*\w*$", "", nm).strip().upper()
        bid = barang_idx.get(base, 0)
        out.append({"barangId": bid, "nama": nm, "qty": it["qty"], "satuan": it["satuan"] or "BUAH",
                    "harga": round(it["harga"], 2), "jenis": "B"})
    return out

CV_MAP = [("PELITA", "PKM"), ("PKM", "PKM"), ("DAKSA", "DKS"), ("DKS", "DKS"), ("ONIX", "OSI"), ("OSI", "OSI")]
def cv_of(label):
    u = label.upper()
    for k, v in CV_MAP:
        if k in u:
            return v
    return ""

# ---------------------------------------------------------------- Digipay 2025
def load_digipay(wb, wbv, by_digi_sheet, barang_idx):
    trx = []
    for ws in wb.worksheets[1:]:
        sat = by_digi_sheet.get(ws.title)
        for b in find_blocks(ws, wbv[ws.title]):
            if not b["items"]:
                continue
            col = b["col"]
            f = {}
            for r in range(max(1, b["hdr"] - 1), min(b["end"], b["hdr"] + 14) + 1):
                k = s(ws.cell(r, col + 7).value)
                if isinstance(k, str) and k and k != "Form Isian Data":
                    f[k.upper().strip()] = ws.cell(r, col + 8).value
            lab = b["label"].upper()
            kanal = "MBIZ" if "MBIZ" in lab else "DIGIPAY"
            tgl = date_from_text(f.get("TANGGAL")) or date_from_text(f.get("TANGGAL PO"))
            t = {
                "cv": "OSI", "kanal": kanal, "satkerId": sat["id"] if sat else 0,
                "tanggal": tgl.isoformat() if tgl else "",
                "noTrx": clean_name(f.get("NO. INV APK") or f.get("NO. PO") or ""),
                "kodeAkun": clean_name(f.get("KODE AKUN ANGGARAN") or ""),
                "noSJ": clean_name(f.get("NO. SURAT JALAN") or ""),
                "items": mk_items(b, barang_idx),
                "inclPPN": kanal == "DIGIPAY" and not b["hasPPN"],
                "status": "lunas" if str(f.get("STATUS", "")).upper() == "PAID" else "selesai",
                "catatan": "Migrasi: REKAP DIGIPAY 2025 / sheet " + ws.title,
            }
            if kanal == "MBIZ":
                t["tipeBayar"] = clean_name(f.get("TYPE PEMBAYARAN") or "")
                t["noBayar"] = clean_name(f.get("NO. PEMBAYARAN") or "")
                tb = date_from_text(f.get("TGL PEMBAYARAN"))
                t["tglBayar"] = tb.isoformat() if tb else ""
                t["wapu"] = "NON" not in str(f.get("TYPE TRANSAKSI", "")).upper()
            trx.append(t)
    return trx

# ---------------------------------------------------------------- BELANJA BPN
MONTH_COL = {chr(ord("J") + i): i + 1 for i in range(12)}

def load_belanja(fn, year, by_kantah_sheet, barang_idx):
    wb = openpyxl.load_workbook(P(fn))
    wbv = openpyxl.load_workbook(P(fn), data_only=True)
    trx = []
    for ws in wb.worksheets[1:]:
        sat = by_kantah_sheet.get(ws.title)
        blocks = find_blocks(ws, wbv[ws.title])
        # pemetaan baris G -> bulan dari rumus J4:U6
        g_month = {}
        for r in (4, 5, 6):
            for colL, mon in MONTH_COL.items():
                v = ws[f"{colL}{r}"].value
                if isinstance(v, str) and v.startswith("="):
                    for g in re.findall(r"G(\d+)", v):
                        g_month[int(g)] = mon
                elif isinstance(v, (int, float)) and v:
                    pass
        for b in blocks:
            if not b["items"]:
                continue
            col = b["col"]
            # label CV: baris atas header
            label = b["label"]
            cv = cv_of(label)
            info = {}
            for r in range(b["hdr"], b["end"] + 1):
                v = ws.cell(r, col + 1).value
                if isinstance(v, str) and ":" in v:
                    k, _, val = v.partition(":")
                    info[k.strip().upper()] = val.strip()
                elif isinstance(v, str) and v.strip().upper() in ("LUNAS",):
                    info["LUNAS"] = True
            mon = None
            for r, _, _ in b["totals"]:
                if r in g_month:
                    mon = g_month[r]
            for it in b["items"]:
                pass
            tsp = date_from_text(info.get("TANGGAL SPK/SP"))
            tbast = date_from_text(info.get("TANGGAL BAST"))
            tkirim = date_from_text(info.get("TANGGAL DI KIRIM"))
            if mon is None:
                ref = tsp or tbast
                mon = ref.month if ref else None
            if tsp and mon and tsp.month == mon and tsp.year == year:
                tgl = tsp
            elif mon:
                tgl = datetime.date(year, mon, 1)
            else:
                tgl = None
            idp = info.get("ID PAKET", "")
            kanal = "INAPROC" if idp.upper().startswith("EP-") else "LAINNYA"
            trx.append({
                "cv": cv, "kanal": kanal, "satkerId": sat["id"] if sat else 0,
                "tanggal": tgl.isoformat() if tgl else "",
                "noTrx": idp, "rup": info.get("NO. RUP", ""), "sumberDana": info.get("SUMBER DANA", ""),
                "tglSP": tsp.isoformat() if tsp else "", "tglBAST": tbast.isoformat() if tbast else "",
                "tglKirim": tkirim.isoformat() if tkirim else "",
                "items": mk_items(b, barang_idx),
                "inclPPN": not b["hasPPN"],
                "status": "lunas" if info.get("LUNAS") else ("selesai" if tbast else "proses"),
                "catatan": f"Migrasi: {fn} / sheet {ws.title}" + (
                    "" if mon else " (bulan tidak tercatat di Excel)") + (
                    "" if idp else (" / " + str(info.get("TANGGAL SPK/SP", "")) if info.get("TANGGAL SPK/SP") and not tsp else "")),
            })
    return trx

def main():
    barang = load_barang()
    barang_idx = {}
    for b in barang:
        barang_idx.setdefault(b["nama"].upper(), b["id"])
    wb_digi = openpyxl.load_workbook(P("DATA REKAP DIGIPAY 2025 CV. ONIX STATIONERY IDEAL.xlsx"))
    satker, by_kantah, by_digi = load_satker(wb_digi)
    by_kantah["KAB. GARUT"]["kodeSatker"] = by_kantah["KAB. GARUT"]["kodeSatker"] or "429734"
    trx = []
    trx += load_digipay(wb_digi, openpyxl.load_workbook(P("DATA REKAP DIGIPAY 2025 CV. ONIX STATIONERY IDEAL.xlsx"), data_only=True), by_digi, barang_idx)
    trx += load_belanja("BELANJA BPN 2025.xlsx", 2025, by_kantah, barang_idx)
    trx += load_belanja("BELANJA BPN 2026.xlsx", 2026, by_kantah, barang_idx)
    for i, t in enumerate(trx, start=1):
        t["id"] = i
    seed = {"barang": barang, "satker": satker, "transaksi": trx}
    json.dump(seed, open(OUT, "w"), ensure_ascii=False, indent=1)
    # ringkasan
    c = collections.Counter((t["tanggal"][:4], t["cv"], t["kanal"]) for t in trx)
    for k in sorted(c):
        print(k, c[k])
    print("barang", len(barang), "satker", len(satker), "transaksi", len(trx))
    print("tanpa tanggal:", sum(1 for t in trx if not t["tanggal"]), "tanpa satker:",
          sum(1 for t in trx if not t["satkerId"]), "tanpa cv:", sum(1 for t in trx if not t["cv"]))

if __name__ == "__main__":
    main()
